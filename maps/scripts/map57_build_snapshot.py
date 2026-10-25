#!/usr/bin/env python3
"""Build Map 57's joined World Bank CPPI and port-coordinate snapshot."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).lower()
    return re.sub(r"[^a-z0-9]+", "", text)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_geojson(path: Path) -> dict:
    # The cataloged file is Latin-1 and currently carries a non-JSON diagnostic
    # line after the FeatureCollection. Keep only the complete JSON object.
    text = path.read_bytes().decode("latin-1")
    end = text.rfind("\n}")
    if end < 0:
        raise ValueError("Port GeoJSON has no complete trailing object")
    return json.loads(text[: end + 2])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("cppi_xlsx", type=Path)
    parser.add_argument("ports_geojson", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    geo = read_geojson(args.ports_geojson)
    by_code = {}
    by_name = defaultdict(list)
    for feature in geo.get("features", []):
        props = feature.get("properties") or {}
        coords = (feature.get("geometry") or {}).get("coordinates")
        if not (isinstance(coords, list) and len(coords) >= 2):
            continue
        record = {
            "locode": str(props.get("LOCODE") or "").strip().upper(),
            "name": str(props.get("Name") or "").strip(),
            "country": str(props.get("Country") or "").strip(),
            "lon": round(float(coords[0]), 6),
            "lat": round(float(coords[1]), 6),
        }
        if record["locode"]:
            by_code[record["locode"]] = record
        by_name[normalized(record["name"])].append(record)

    sheet = load_workbook(args.cppi_xlsx, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}
    required = ["Port", "LOCODE", "Territory", "World Bank Region", "World Bank Income Group",
                "Maritime services region", "CPPI 2020", "CPPI 2021", "CPPI 2022", "CPPI 2023",
                "CPPI 2024", "Rank", "Berth hours in % of port hours", "Statistical Index",
                "Administrative Index"]
    missing = [name for name in required if name not in positions]
    if missing:
        raise ValueError(f"Missing CPPI columns: {missing}")

    output = []
    methods = Counter()
    unmatched = []
    for values in rows:
        get = lambda name: values[positions[name]]
        locode = str(get("LOCODE") or "").strip().upper()
        point = by_code.get(locode)
        method = "locode"
        if point is None:
            candidates = by_name.get(normalized(get("Port")), [])
            territory = normalized(get("Territory"))
            country_matches = [item for item in candidates if normalized(item["country"]) == territory]
            if len(country_matches) == 1:
                point = country_matches[0]
                method = "unique_name_country"
            elif len(candidates) == 1:
                point = candidates[0]
                method = "unique_name"
        if point is None:
            unmatched.append({"port": get("Port"), "locode": locode, "territory": get("Territory")})
            continue
        methods[method] += 1
        scores = {str(year): get(f"CPPI {year}") for year in range(2020, 2025)}
        output.append({
            "port": get("Port"), "locode": locode, "territory": get("Territory"),
            "region": get("World Bank Region"), "income": get("World Bank Income Group"),
            "maritime_region": get("Maritime services region"), "lat": point["lat"], "lon": point["lon"],
            "scores": scores, "rank_2024": get("Rank"), "berth_share": get("Berth hours in % of port hours"),
            "statistical_index_2024": get("Statistical Index"),
            "administrative_index_2024": get("Administrative Index"), "coordinate_match": method,
        })

    output.sort(key=lambda item: (item["port"], item["locode"]))
    payload = {
        "schema": "map57-world-bank-cppi-v1",
        "built_on": "2026-10-25",
        "cppi_vintage": "2020-2024",
        "coordinate_vintage": "2020-07-10",
        "source_rows": len(output) + len(unmatched),
        "mapped_rows": len(output),
        "unmatched_rows": len(unmatched),
        "match_methods": dict(sorted(methods.items())),
        "ports": output,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(json.dumps({
        "output": str(args.output), "output_bytes": args.output.stat().st_size,
        "output_sha256": sha256(args.output), "cppi_sha256": sha256(args.cppi_xlsx),
        "ports_geojson_sha256": sha256(args.ports_geojson), **{key: payload[key] for key in
        ("source_rows", "mapped_rows", "unmatched_rows", "match_methods")},
        "unmatched": unmatched,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
