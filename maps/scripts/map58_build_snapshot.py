#!/usr/bin/env python3
"""Build Map 58's normalized FAO AQUASTAT dam-point snapshot."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def clean(value: object) -> str:
    return " ".join(str(value or "").split())


def number(value: object):
    try:
        result = float(value)
        return round(result, 6) if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def year(value: object):
    result = number(value)
    return int(result) if result is not None and 1000 <= result <= 2200 else None


def flag(value: object) -> bool:
    return clean(value).lower() not in {"", "0", "false", "no", "none", "-"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument("workbooks", nargs="+", type=Path)
    args = parser.parse_args()

    dams, skipped, regions = [], Counter(), Counter()
    input_hashes = {}
    for path in sorted(args.workbooks, key=lambda p: p.name):
        region = path.stem.replace("-", " ")
        input_hashes[path.name] = sha256(path)
        sheet = load_workbook(path, read_only=True, data_only=True)["Dams"]
        rows = sheet.iter_rows(values_only=True)
        next(rows)  # title row
        headers = [clean(value) for value in next(rows)]
        pos = {name: index for index, name in enumerate(headers)}
        required = ["Country", "Name of dam", "Decimal degree latitude", "Decimal degree longitude"]
        if any(name not in pos for name in required):
            raise ValueError(f"{path}: missing required columns")
        for values in rows:
            get = lambda name: values[pos[name]] if pos[name] < len(values) else None
            name, country = clean(get("Name of dam")), clean(get("Country"))
            if not name and not country:
                continue
            lat, lon = number(get("Decimal degree latitude")), number(get("Decimal degree longitude"))
            if lat is None or lon is None:
                skipped["missing_coordinates"] += 1
                continue
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                skipped["invalid_coordinates"] += 1
                continue
            purposes = []
            for source, label in (("Irrigation", "Irrigation"), ("Water supply", "Water supply"),
                                  ("Flood control", "Flood control"), ("Hydroelectricity (MW)", "Hydropower"),
                                  ("Navigation", "Navigation"), ("Recreation", "Recreation"),
                                  ("Pollution control", "Pollution control"), ("Livestock rearing", "Livestock")):
                if source in pos and flag(get(source)):
                    purposes.append(label)
            dams.append({
                "name": name or "Unnamed dam", "country": country, "iso3": clean(get("ISO alpha- 3")),
                "admin": clean(get("Administrative Unit")), "city": clean(get("Nearest city")),
                "river": clean(get("River")), "basin": clean(get("Major basin")),
                "year": year(get("Completed /operational since")), "height_m": number(get("Dam height (m)")),
                "capacity_mcm": number(get("Reservoir capacity (million m3)")),
                "area_km2": number(get("Reservoir area (km2)")), "purposes": purposes,
                "lat": lat, "lon": lon, "region": region,
            })
            regions[region] += 1

    dams.sort(key=lambda item: (item["country"], item["name"], item["lat"], item["lon"]))
    payload = {
        "schema": "map58-fao-aquastat-dams-v1", "built_on": "2026-10-26",
        "catalog_edition": "2026", "catalog_revision": "2026-07-16",
        "mapped_records": len(dams), "omitted_records": sum(skipped.values()),
        "omission_reasons": dict(sorted(skipped.items())), "records_by_region": dict(sorted(regions.items())),
        "dams": dams,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    serialized = (json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n").encode()
    if args.output.suffix == ".gz":
        args.output.write_bytes(gzip.compress(serialized, compresslevel=9, mtime=0))
    else:
        args.output.write_bytes(serialized)
    print(json.dumps({"output": str(args.output), "output_bytes": args.output.stat().st_size,
                      "output_sha256": sha256(args.output), "input_sha256": input_hashes,
                      "mapped_records": len(dams), "omitted_records": sum(skipped.values()),
                      "omission_reasons": dict(skipped), "records_by_region": dict(regions)}, indent=2))


if __name__ == "__main__":
    main()
